import io
import csv
import json
import zipfile
import logging
from datetime import datetime, timedelta

from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse, JsonResponse
from django.db.models import Q
from django.conf import settings

from experiments.models import Experiment
from reports.models import Report
from audit.models import AuditLog

export_logger = logging.getLogger("reports")


@login_required
def export_experiments(request):
    """Bulk export user's experiments as CSV, JSON, or XLSX with filters."""
    user = request.user
    is_admin = user.is_staff or user.user_type == "admin"

    experiments = Experiment.objects.all() if is_admin else Experiment.objects.filter(user=user)

    # Apply filters
    technique = request.GET.get("technique")
    date_from = request.GET.get("date_from")
    date_to = request.GET.get("date_to")
    status = request.GET.get("status")
    fmt = request.GET.get("format", "csv")

    if technique:
        experiments = experiments.filter(privacy_technique__technique_type=technique)
    if date_from:
        experiments = experiments.filter(created_at__gte=date_from)
    if date_to:
        experiments = experiments.filter(created_at__lte=date_to + "T23:59:59")
    if status:
        experiments = experiments.filter(status=status)

    experiments = experiments.select_related(
        "privacy_technique", "dataset", "user"
    ).order_by("-created_at")

    # If it's just a page view (no download trigger)
    if request.GET.get("action") != "download":
        from privacy_tools.models import PrivacyTechnique
        techniques = PrivacyTechnique.objects.all()
        return render(request, "export/export_page.html", {
            "experiments": experiments,
            "techniques": techniques,
            "current_filters": {
                "technique": technique,
                "date_from": date_from,
                "date_to": date_to,
                "status": status,
                "format": fmt,
            },
            "export_type": "experiments",
        })

    # Build filename base
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    base = f"experiments_export_{ts}"

    if fmt == "zip":
        return _zip_experiments(experiments, base)
    elif fmt == "xlsx":
        return _xlsx_experiments(experiments, base)
    elif fmt == "json":
        return _json_experiments(experiments, base)
    else:
        return _csv_experiments(experiments, base)


@login_required
def export_reports(request):
    """Bulk export report metadata as CSV or JSON."""
    user = request.user
    is_admin = user.is_staff or user.user_type == "admin"

    reports = Report.objects.all() if is_admin else Report.objects.filter(user=user)

    date_from = request.GET.get("date_from")
    date_to = request.GET.get("date_to")
    report_type = request.GET.get("report_type")
    fmt = request.GET.get("format", "csv")

    if date_from:
        reports = reports.filter(created_at__gte=date_from)
    if date_to:
        reports = reports.filter(created_at__lte=date_to + "T23:59:59")
    if report_type:
        reports = reports.filter(report_type=report_type)

    reports = reports.select_related("user").order_by("-created_at")

    if request.GET.get("action") != "download":
        return render(request, "export/export_page.html", {
            "reports": reports,
            "current_filters": {
                "date_from": date_from,
                "date_to": date_to,
                "report_type": report_type,
                "format": fmt,
            },
            "export_type": "reports",
        })

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    base = f"reports_export_{ts}"

    if fmt == "json":
        return _json_reports(reports, base)
    elif fmt == "zip":
        return _zip_reports(reports, base)
    else:
        return _csv_reports(reports, base)


# ─── Experiment Export Helpers ───────────────────────────────


def _csv_experiments(qs, base):
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([
        "ID", "Name", "Technique", "Dataset", "Status",
        "Privacy Score", "Accuracy", "Execution Time (s)",
        "Throughput (tps)", "Anonymity Set Size", "Created At",
    ])
    for exp in qs:
        writer.writerow([
            exp.id, exp.name, exp.privacy_technique.name, exp.dataset.name,
            exp.get_status_display(), exp.privacy_score or "",
            exp.accuracy or "", exp.execution_time or "",
            exp.throughput or "", exp.anonymity_set_size or "",
            exp.created_at.strftime("%Y-%m-%d %H:%M:%S") if exp.created_at else "",
        ])
    response = HttpResponse(buf.getvalue(), content_type="text/csv")
    response["Content-Disposition"] = f'attachment; filename="{base}.csv"'
    _log_export(request=None, count=qs.count(), export_type="experiments", fmt="csv")
    return response


def _json_experiments(qs, base):
    data = []
    for exp in qs:
        data.append({
            "id": exp.id,
            "name": exp.name,
            "technique": exp.privacy_technique.name,
            "technique_type": exp.privacy_technique.technique_type,
            "dataset": exp.dataset.name,
            "status": exp.status,
            "privacy_score": exp.privacy_score,
            "accuracy": exp.accuracy,
            "execution_time": exp.execution_time,
            "throughput": exp.throughput,
            "anonymity_set_size": exp.anonymity_set_size,
            "configuration": exp.configuration,
            "metrics": exp.metrics,
            "error_message": exp.error_message,
            "created_at": exp.created_at.isoformat() if exp.created_at else None,
            "completed_at": exp.completed_at.isoformat() if exp.completed_at else None,
        })
    response = HttpResponse(
        json.dumps(data, indent=2, default=str),
        content_type="application/json",
    )
    response["Content-Disposition"] = f'attachment; filename="{base}.json"'
    _log_export(request=None, count=len(data), export_type="experiments", fmt="json")
    return response


def _xlsx_experiments(qs, base):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Experiments"
    headers = [
        "ID", "Name", "Technique", "Dataset", "Status",
        "Privacy Score", "Accuracy", "Execution Time (s)",
        "Throughput (tps)", "Anonymity Set Size", "Configuration",
    ]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="8B5CF6", end_color="8B5CF6", fill_type="solid")
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
    for row, exp in enumerate(qs, 2):
        ws.cell(row=row, column=1, value=exp.id)
        ws.cell(row=row, column=2, value=exp.name)
        ws.cell(row=row, column=3, value=exp.privacy_technique.name)
        ws.cell(row=row, column=4, value=exp.dataset.name)
        ws.cell(row=row, column=5, value=exp.get_status_display())
        ws.cell(row=row, column=6, value=exp.privacy_score or "")
        ws.cell(row=row, column=7, value=exp.accuracy or "")
        ws.cell(row=row, column=8, value=exp.execution_time or "")
        ws.cell(row=row, column=9, value=exp.throughput or "")
        ws.cell(row=row, column=10, value=exp.anonymity_set_size or "")
        ws.cell(row=row, column=11, value=json.dumps(exp.configuration) if exp.configuration else "")
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[chr(64 + col)].width = 20
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    response = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{base}.xlsx"'
    _log_export(request=None, count=qs.count(), export_type="experiments", fmt="xlsx")
    return response


def _zip_experiments(qs, base):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        # CSV
        csv_buf = io.StringIO()
        writer = csv.writer(csv_buf)
        writer.writerow([
            "ID", "Name", "Technique", "Dataset", "Status",
            "Privacy Score", "Accuracy", "Execution Time (s)",
            "Throughput", "Anonymity Set Size", "Created At",
        ])
        for exp in qs:
            writer.writerow([
                exp.id, exp.name, exp.privacy_technique.name, exp.dataset.name,
                exp.get_status_display(), exp.privacy_score or "",
                exp.accuracy or "", exp.execution_time or "",
                exp.throughput or "", exp.anonymity_set_size or "",
                exp.created_at.strftime("%Y-%m-%d %H:%M:%S") if exp.created_at else "",
            ])
        zf.writestr(f"{base}.csv", csv_buf.getvalue())
        # JSON
        data = []
        for exp in qs:
            data.append({
                "id": exp.id, "name": exp.name,
                "technique": exp.privacy_technique.name,
                "status": exp.status,
                "privacy_score": exp.privacy_score,
                "accuracy": exp.accuracy,
                "execution_time": exp.execution_time,
                "throughput": exp.throughput,
                "configuration": exp.configuration,
                "metrics": exp.metrics,
            })
        zf.writestr(f"{base}.json", json.dumps(data, indent=2, default=str))
    buf.seek(0)
    response = HttpResponse(buf.getvalue(), content_type="application/zip")
    response["Content-Disposition"] = f'attachment; filename="{base}.zip"'
    _log_export(request=None, count=qs.count(), export_type="experiments", fmt="zip")
    return response


# ─── Report Export Helpers ───────────────────────────────────


def _csv_reports(qs, base):
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([
        "ID", "Title", "Type", "Format", "Experiment Count",
        "Created At", "Generated At",
    ])
    for r in qs:
        writer.writerow([
            r.id, r.title, r.get_report_type_display(), r.file_format,
            r.experiments.count(),
            r.created_at.strftime("%Y-%m-%d %H:%M:%S") if r.created_at else "",
            r.generated_at.strftime("%Y-%m-%d %H:%M:%S") if r.generated_at else "",
        ])
    response = HttpResponse(buf.getvalue(), content_type="text/csv")
    response["Content-Disposition"] = f'attachment; filename="{base}.csv"'
    return response


def _json_reports(qs, base):
    data = []
    for r in qs:
        data.append({
            "id": r.id,
            "title": r.title,
            "report_type": r.report_type,
            "file_format": r.file_format,
            "experiment_count": r.experiments.count(),
            "user": r.user.username,
            "created_at": r.created_at.isoformat() if r.created_at else None,
            "generated_at": r.generated_at.isoformat() if r.generated_at else None,
        })
    response = HttpResponse(
        json.dumps(data, indent=2, default=str),
        content_type="application/json",
    )
    response["Content-Disposition"] = f'attachment; filename="{base}.json"'
    return response


def _zip_reports(qs, base):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        csv_buf = io.StringIO()
        writer = csv.writer(csv_buf)
        writer.writerow([
            "ID", "Title", "Type", "Format", "Experiment Count", "Created At",
        ])
        for r in qs:
            writer.writerow([
                r.id, r.title, r.get_report_type_display(), r.file_format,
                r.experiments.count(),
                r.created_at.strftime("%Y-%m-%d %H:%M:%S") if r.created_at else "",
            ])
        zf.writestr(f"{base}.csv", csv_buf.getvalue())
        json_data = [
            {
                "id": r.id, "title": r.title, "report_type": r.report_type,
                "file_format": r.file_format,
                "experiment_count": r.experiments.count(),
                "created_at": r.created_at.isoformat() if r.created_at else None,
            }
            for r in qs
        ]
        zf.writestr(f"{base}.json", json.dumps(json_data, indent=2, default=str))
    buf.seek(0)
    response = HttpResponse(buf.getvalue(), content_type="application/zip")
    response["Content-Disposition"] = f'attachment; filename="{base}.zip"'
    return response


# ─── Shared Helpers ──────────────────────────────────────────


def _log_export(request, count, export_type, fmt):
    try:
        AuditLog.objects.create(
            action_type="EXPORT",
            content_type=export_type.title(),
            object_id=0,
            object_repr=f"Bulk export of {count} {export_type} as {fmt}",
            changes={"count": count, "format": fmt},
        )
    except Exception:
        pass
